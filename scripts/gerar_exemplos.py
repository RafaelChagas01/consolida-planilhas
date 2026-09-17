"""Gera as planilhas de exemplo em exemplos/. Os dados sao ficticios.

Cada loja exporta de um jeito, como costuma acontecer quando cada gerente monta a propria planilha.

Uso: python scripts/gerar_exemplos.py
"""

import random
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook

PASTA = Path("exemplos")

PRODUTOS = [
    ("Cimento CP II 50kg", "Cimento e argamassa", 36.90),
    ("Argamassa AC-I 20kg", "Cimento e argamassa", 19.90),
    ("Cal hidratada 20kg", "Cimento e argamassa", 17.50),
    ("Areia média saco 20kg", "Cimento e argamassa", 6.90),
    ("Tubo PVC 25mm 6m", "Hidráulica", 24.90),
    ("Joelho PVC 25mm", "Hidráulica", 1.90),
    ("Registro de gaveta 3/4", "Hidráulica", 58.90),
    ("Caixa d'água 500L", "Hidráulica", 389.00),
    ("Fio flexível 2,5mm 100m", "Elétrica", 219.00),
    ("Disjuntor 20A", "Elétrica", 18.90),
    ("Tomada 10A", "Elétrica", 9.90),
    ("Lâmpada LED 9W", "Elétrica", 8.50),
    ("Martelo unha 27mm", "Ferramentas", 29.90),
    ("Trena 5m", "Ferramentas", 24.90),
    ("Furadeira 650W", "Ferramentas", 229.00),
    ("Jogo de chaves de fenda", "Ferramentas", 39.90),
    ("Tinta acrílica 18L branco", "Tintas", 289.00),
    ("Rolo de lã 23cm", "Tintas", 22.90),
    ("Massa corrida 25kg", "Tintas", 64.90),
    ("Lixa grão 120", "Tintas", 1.50),
    ("Porcelanato 60x60 (m²)", "Pisos e revestimentos", 69.90),
    ("Rejunte 1kg", "Pisos e revestimentos", 7.90),
]

SEM_ACENTO = str.maketrans("ÁÂÃÉÊÍÓÔÕÚÇ", "AAAEEIOOOUC")


def vendas(rng: random.Random, quantidade: int) -> list[tuple[date, str, str, int, float]]:
    inicio = date(2026, 1, 2)
    dias = sorted(inicio + timedelta(days=rng.randrange(89)) for _ in range(quantidade))
    linhas = []
    for dia in dias:
        if dia.weekday() == 6:
            dia -= timedelta(days=1)
        while True:
            produto, categoria, preco = rng.choice(PRODUTOS)
            qtd = rng.choice([1, 1, 1, 2, 2, 3, 5, 10]) if preco < 100 else rng.choice([1, 1, 2])
            # as unicas linhas repetidas devem ser as coladas de proposito
            if not linhas or linhas[-1] != (dia, produto, categoria, qtd, preco):
                break
        linhas.append((dia, produto, categoria, qtd, preco))
    return linhas


def loja_centro(rng: random.Random) -> None:
    # cabecalho na primeira linha, datas de verdade, linha de total no fim
    wb = Workbook()
    ws = wb.active
    ws.title = "Vendas"
    ws.append(["Data", "Produto", "Categoria", "Qtd", "Valor Unit.", "Total"])
    linhas = vendas(rng, 520)
    for i, (dia, produto, categoria, qtd, preco) in enumerate(linhas):
        total = round(qtd * preco, 2)
        if i == 211:
            total = round(total * 10, 2)
        ws.append([dia, produto, categoria, qtd, preco, total])
        if i in (80, 81, 300):
            ws.append([])
    ws.append([None, "TOTAL", None, None, None, round(sum(q * p for *_, q, p in linhas), 2)])
    for cell in ws["A"][1:]:
        cell.number_format = "DD/MM/YYYY"
    wb.save(PASTA / "loja_centro.xlsx")


def loja_rudge_ramos(rng: random.Random) -> None:
    # titulo antes do cabecalho, tudo digitado como texto
    wb = Workbook()
    ws = wb.active
    ws.title = "Planilha1"
    ws.append(["Relatório de vendas - 1º trimestre 2026"])
    ws.append(["Loja Rudge Ramos"])
    ws.append([])
    ws.append(["data da venda", "descrição", "categoria", "quantidade", "preço"])
    for i, (dia, produto, categoria, qtd, preco) in enumerate(vendas(rng, 430)):
        texto_data = dia.strftime("%d/%m/%Y")
        texto_preco = f"R$ {preco:,.2f}".replace(",", "_").replace(".", ",").replace("_", ".")
        if i == 57:
            texto_data = "31/02/2026"
        if i == 190:
            qtd = None
        ws.append([texto_data, produto, categoria, None if qtd is None else str(qtd), texto_preco])
    wb.save(PASTA / "vendas_rudge_ramos.xlsx")


def loja_diadema(rng: random.Random) -> None:
    # nomes de coluna diferentes, categoria em maiusculo sem acento, linhas coladas duas vezes e devolucao
    wb = Workbook()
    ws = wb.active
    ws.title = "JAN-MAR"
    ws.append(["DATA", "PRODUTO", "GRUPO", "QTDE", "VL UNIT", "VL TOTAL"])
    linhas = vendas(rng, 360)
    for i, (dia, produto, categoria, qtd, preco) in enumerate(linhas):
        grupo = categoria.upper().translate(SEM_ACENTO)
        if i % 7 == 0:
            grupo += " "
        linha = [dia, produto.upper(), grupo, qtd, preco, round(qtd * preco, 2)]
        ws.append(linha)
        if i in (40, 41, 42):
            ws.append(linha)
        if i == 150:
            ws.append([dia, produto.upper(), grupo, -1, preco, -preco])
    for cell in ws["A"][1:]:
        cell.number_format = "DD/MM/YY"
    wb.save(PASTA / "Loja Diadema - jan a mar.xlsx")


def main() -> None:
    PASTA.mkdir(exist_ok=True)
    rng = random.Random(2026)
    loja_centro(rng)
    loja_rudge_ramos(rng)
    loja_diadema(rng)
    print("planilhas geradas em", PASTA)


if __name__ == "__main__":
    main()
