from datetime import date
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook

from consolida.arquivo import ArquivoInvalido
from consolida.processar import processar


@pytest.fixture
def relatorio():
    livro = Workbook()
    aba = livro.active
    aba.append(["Data", "Produto", "Categoria", "Qtd", "Preço", "Total"])
    aba.append([date(2026, 1, 5), None, "Ferramentas", 1, 10, 10])
    aba.append([date(2026, 1, 6), "Trena 5m", "Ferramentas", 2, 24.9, 49.8])
    aba.append([date(2026, 2, 6), "Trena 5m", "Ferramentas", 1, 24.9, 99])
    # celula de texto que comeca com "=", como fica quando alguem cola de um CSV
    aba["B2"].value = '=HYPERLINK("http://exemplo.invalid","clique")'
    aba["B2"].data_type = "s"
    saida = BytesIO()
    livro.save(saida)

    resultado = processar([("loja_centro.xlsx", saida.getvalue())])
    return load_workbook(BytesIO(resultado.planilha))


def test_texto_da_planilha_enviada_nao_vira_formula(relatorio):
    vendas = relatorio["Vendas"]
    produtos = [vendas.cell(row=i, column=3) for i in range(2, vendas.max_row + 1)]
    injetado = next(c for c in produtos if c.value.startswith("=HYPERLINK"))
    assert injetado.data_type == "s"
    assert injetado.quotePrefix


def test_resumo_usa_formulas_nos_totais(relatorio):
    resumo = relatorio["Resumo"]
    valores = [[c.value for c in linha] for linha in resumo.iter_rows(max_col=5)]
    total = next(linha for linha in valores if linha[0] == "Total")
    assert total[1:] == ["=SUM(B8:B8)", "=SUM(C8:C8)", "=SUM(D8:D8)", "=D9/D$9"]
    assert "1 linhas ficaram de fora. O motivo de cada uma está na aba Descartadas." in [linha[0] for linha in valores]


def test_abas_de_vendas_e_descartadas_viram_tabela(relatorio):
    assert "Vendas" in relatorio["Vendas"].tables
    assert relatorio["Vendas"].freeze_panes == "A2"
    descartadas = relatorio["Descartadas"]
    assert "Descartadas" in descartadas.tables
    assert descartadas["C2"].value.startswith("total informado")


def test_sem_nenhuma_venda_valida(xlsx):
    conteudo = xlsx([["Data", "Produto", "Qtd", "Preço"], ["ontem", "Lixa", 1, 1.5]])
    with pytest.raises(ArquivoInvalido, match="nenhuma linha de venda válida"):
        processar([("a.xlsx", conteudo)])


def test_limite_de_arquivos(xlsx):
    conteudo = xlsx([["Data", "Produto", "Qtd", "Preço"], ["02/01/2026", "Lixa", 1, 1.5]])
    with pytest.raises(ArquivoInvalido, match="no máximo 5"):
        processar([(f"{i}.xlsx", conteudo) for i in range(6)])


def test_erro_mostra_qual_arquivo(xlsx):
    with pytest.raises(ArquivoInvalido, match=r"^vendas\.csv: só aceito"):
        processar([("../vendas.csv", b"data;produto")])
