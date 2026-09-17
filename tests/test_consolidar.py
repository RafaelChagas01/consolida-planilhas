from datetime import date
from io import BytesIO

from openpyxl import load_workbook

from consolida.consolidar import consolidar
from consolida.leitura import ArquivoLido
from consolida.processar import processar


def venda(linha, produto="Trena 5m", categoria="Ferramentas", dia=date(2026, 1, 5), quantidade=1.0, preco=24.9, loja="Centro"):
    return {
        "loja": loja,
        "data": dia,
        "produto": produto,
        "categoria": categoria,
        "quantidade": quantidade,
        "preco": preco,
        "total": round(quantidade * preco, 2),
        "arquivo": f"{loja}.xlsx",
        "linha": linha,
    }


def test_junta_nomes_escritos_de_jeitos_diferentes():
    lidos = [
        ArquivoLido("centro.xlsx", "Centro", [venda(2, categoria="Hidráulica"), venda(3, categoria="Hidráulica ")]),
        ArquivoLido("diadema.xlsx", "Diadema", [venda(2, produto="TRENA 5M", categoria="HIDRAULICA", loja="Diadema")]),
    ]
    dados = consolidar(lidos)
    assert set(dados.vendas["categoria"]) == {"Hidráulica"}
    assert set(dados.vendas["produto"]) == {"Trena 5m"}


def test_nome_so_em_maiusculo_fica_legivel():
    dados = consolidar([ArquivoLido("a.xlsx", "Diadema", [venda(2, categoria="ELETRICA")])])
    assert dados.vendas["categoria"].tolist() == ["Eletrica"]


def test_marca_so_a_linha_identica_a_de_cima():
    vendas = [venda(2), venda(3), venda(4, produto="Lixa"), venda(5)]
    dados = consolidar([ArquivoLido("a.xlsx", "Centro", vendas)])
    marcadas = dados.vendas.loc[dados.vendas["repetida"], "linha"].tolist()
    assert marcadas == [3]
    assert dados.repetidas == 1
    assert len(dados.vendas) == 4


def test_totais_por_loja_mes_e_categoria():
    lidos = [
        ArquivoLido(
            "centro.xlsx",
            "Centro",
            [venda(2, quantidade=2), venda(3, produto="Tinta 18L", categoria="Tintas", preco=289.0, dia=date(2026, 2, 1))],
        ),
        ArquivoLido("diadema.xlsx", "Diadema", [venda(2, loja="Diadema", quantidade=4)]),
    ]
    dados = consolidar(lidos)

    assert dados.por_loja.to_dict("records") == [
        {"loja": "Centro", "vendas": 2, "itens": 3.0, "faturamento": 338.8},
        {"loja": "Diadema", "vendas": 1, "itens": 4.0, "faturamento": 99.6},
    ]
    assert dados.por_mes.loc["2026-01-01"].to_dict() == {"Centro": 49.8, "Diadema": 99.6}
    assert dados.por_categoria["categoria"].tolist() == ["Tintas", "Ferramentas"]
    assert dados.top_produtos.iloc[0]["produto"] == "Tinta 18L"


def test_processa_as_planilhas_de_exemplo():
    from pathlib import Path

    arquivos = [(p.name, p.read_bytes()) for p in sorted(Path("exemplos").glob("*.xlsx"))]
    resultado = processar(arquivos)

    assert [a["loja"] for a in resultado.dados.arquivos] == ["Diadema", "Centro", "Rudge Ramos"]
    assert len(resultado.dados.descartadas) == 5
    assert resultado.dados.repetidas == 3

    livro = load_workbook(BytesIO(resultado.planilha))
    assert livro.sheetnames == ["Resumo", "Vendas", "Descartadas"]
    assert livro["Vendas"].max_row == len(resultado.dados.vendas) + 1
