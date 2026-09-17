import base64
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import app

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@pytest.fixture
def client():
    return TestClient(app)


@pytest.fixture
def planilha_ok(xlsx):
    return xlsx([["Data", "Produto", "Qtd", "Preço"], ["02/01/2026", "Lixa", 10, "1,50"], ["03/01/2026", "Trena", 0, 24.9]])


def test_exemplo(client):
    resposta = client.post("/api/exemplo")
    assert resposta.status_code == 200
    dados = resposta.json()
    assert dados["periodo"] == {"inicio": "02/01/2026", "fim": "31/03/2026"}
    assert [loja["loja"] for loja in dados["lojas"]] == ["Centro", "Rudge Ramos", "Diadema"]
    assert dados["total_descartadas"] == 5
    assert dados["repetidas"] == 3
    assert load_workbook(BytesIO(base64.b64decode(dados["planilha"]))).sheetnames == ["Resumo", "Vendas", "Descartadas"]
    assert resposta.headers["cache-control"] == "no-store"


def test_envio(client, planilha_ok):
    resposta = client.post("/api/consolidar", files=[("arquivos", ("loja_centro.xlsx", planilha_ok, XLSX))])
    assert resposta.status_code == 200
    dados = resposta.json()
    assert dados["vendas"] == 1
    assert dados["faturamento"] == 15.0
    assert dados["descartadas"] == [
        {"arquivo": "loja_centro.xlsx", "linha": 3, "motivo": "quantidade zero ou negativa (devolução não entra no fechamento)"}
    ]
    assert "conteudo" not in dados["descartadas"][0]


def test_arquivo_invalido(client):
    resposta = client.post("/api/consolidar", files=[("arquivos", ("vendas.xlsx", b"nao e planilha", XLSX))])
    assert resposta.status_code == 422
    assert resposta.json() == {"erro": "vendas.xlsx: não é um .xlsx válido (pode estar protegido por senha)"}


def test_sem_arquivo(client):
    resposta = client.post("/api/consolidar", data={"site": ""})
    assert resposta.status_code == 422


def test_mais_de_cinco_arquivos(client, planilha_ok):
    arquivos = [("arquivos", (f"{i}.xlsx", planilha_ok, XLSX)) for i in range(6)]
    resposta = client.post("/api/consolidar", files=arquivos)
    assert resposta.status_code == 400


def test_campo_escondido_preenchido(client, planilha_ok):
    resposta = client.post("/api/consolidar", data={"site": "http://spam.invalid"}, files=[("arquivos", ("a.xlsx", planilha_ok, XLSX))])
    assert resposta.status_code == 400


def test_corpo_grande_demais(client):
    resposta = client.post("/api/consolidar", content=b"x", headers={"content-length": str(5 * 1024 * 1024)})
    assert resposta.status_code == 413


def test_corpo_grande_demais_sem_content_length(client):
    def partes():
        for _ in range(5):
            yield b"0" * (1024 * 1024)

    resposta = client.post("/api/consolidar", content=partes(), headers={"content-type": "multipart/form-data; boundary=x"})
    assert "content-length" not in resposta.request.headers
    assert resposta.status_code == 413


def test_exemplo_sem_content_length(client):
    # na Vercel o POST sem corpo chega sem esse cabecalho
    resposta = client.post("/api/exemplo", content=iter([b""]))
    assert "content-length" not in resposta.request.headers
    assert resposta.status_code == 200


def test_limite_por_ip(client):
    codigos = [client.post("/api/consolidar", data={"site": ""}).status_code for _ in range(13)]
    assert codigos[:12] == [422] * 12
    assert codigos[12] == 429


def test_download_de_exemplo_so_da_lista(client):
    assert client.get("/api/exemplos/loja_centro.xlsx").status_code == 200
    assert client.get("/api/exemplos/..%2Fapp%2Fmain.py").status_code == 404
    assert client.get("/api/exemplos/nao-existe.xlsx").status_code == 404


def test_cabecalhos_de_seguranca(client):
    resposta = client.get("/api/health")
    assert "frame-ancestors 'none'" in resposta.headers["content-security-policy"]
    assert resposta.headers["x-content-type-options"] == "nosniff"
    assert resposta.headers["x-frame-options"] == "DENY"
