import re
from dataclasses import dataclass, field
from datetime import date
from io import BytesIO
from pathlib import PurePath

from openpyxl import load_workbook

from consolida.arquivo import ArquivoInvalido, validar_xlsx
from consolida.colunas import achar_cabecalho
from consolida.conversao import chave, data, numero, texto, vazio

LINHAS_MAXIMAS = 50_000
COLUNAS_LIDAS = 40
ABAS_MAXIMAS = 10
TOLERANCIA_TOTAL = 0.05

MESES = {
    "jan", "fev", "mar", "abr", "mai", "jun", "jul", "ago", "set", "out", "nov", "dez",
    "janeiro", "fevereiro", "marco", "abril", "maio", "junho", "julho", "agosto", "setembro", "outubro", "novembro", "dezembro",
}  # fmt: skip
PALAVRAS_IGNORADAS = MESES | {"loja", "lojas", "vendas", "venda", "filial", "planilha", "relatorio", "fechamento", "trimestre", "tri"}
CONECTORES = {"a", "e", "de", "da", "do", "das", "dos", "ate"}


@dataclass
class ArquivoLido:
    arquivo: str
    loja: str
    vendas: list[dict] = field(default_factory=list)
    descartadas: list[dict] = field(default_factory=list)


def nome_da_loja(arquivo: str, posicao: int) -> str:
    # "Loja Diadema - jan a mar.xlsx" -> "Diadema"
    palavras = [p for p in re.split(r"[^\w]+|_", PurePath(arquivo).stem) if p]
    palavras = [p for p in palavras if chave(p) not in PALAVRAS_IGNORADAS and not any(c.isdigit() for c in p)]
    while palavras and chave(palavras[0]) in CONECTORES:
        palavras.pop(0)
    while palavras and chave(palavras[-1]) in CONECTORES:
        palavras.pop()
    nome = " ".join(p.lower() if chave(p) in CONECTORES else p.capitalize() for p in palavras)[:40].strip()
    return nome or f"Arquivo {posicao}"


def nome_seguro(arquivo: str) -> str:
    base = PurePath(arquivo.replace("\\", "/")).name
    return re.sub(r"[^\w .()-]", "", base)[:80] or "planilha.xlsx"


def _formatar(valor: float) -> str:
    return f"{valor:,.2f}".replace(",", "_").replace(".", ",").replace("_", ".")


def _conteudo(linha: tuple) -> str:
    partes = [v.strftime("%d/%m/%Y") if isinstance(v, date) else texto(v) for v in linha if not vazio(v)]
    return " | ".join(partes)[:200]


def ler_linha(linha: tuple, mapa: dict[str, int]) -> tuple[dict | None, str | None]:
    """Devolve (venda, None) ou (None, motivo do descarte)."""

    def campo(nome):
        indice = mapa.get(nome)
        return linha[indice] if indice is not None and indice < len(linha) else None

    produto = texto(campo("produto"))
    dia = data(campo("data"))
    if dia is None and chave(produto).startswith(("total", "subtotal")):
        return None, "linha de total da planilha"
    if dia is None:
        return None, "data vazia ou inválida"
    if not 2000 <= dia.year <= 2100:
        return None, "data fora do intervalo esperado"
    if not produto:
        return None, "produto vazio"

    quantidade = numero(campo("quantidade"))
    if quantidade is None:
        return None, "quantidade vazia ou inválida"
    if quantidade <= 0:
        return None, "quantidade zero ou negativa (devolução não entra no fechamento)"

    preco = numero(campo("preco"))
    if preco is None or preco < 0:
        return None, "preço vazio ou inválido"

    calculado = round(quantidade * preco, 2)
    informado = numero(campo("total"))
    if informado is not None and abs(informado - calculado) > TOLERANCIA_TOTAL:
        return None, f"total informado ({_formatar(informado)}) diferente de quantidade x preço ({_formatar(calculado)})"

    return {
        "loja": texto(campo("loja")),
        "data": dia,
        "produto": produto,
        "categoria": texto(campo("categoria")) or "Sem categoria",
        "quantidade": quantidade,
        "preco": preco,
        "total": calculado,
    }, None


def ler_planilha(arquivo: str, conteudo: bytes, posicao: int = 1) -> ArquivoLido:
    validar_xlsx(arquivo, conteudo)
    nome = nome_seguro(arquivo)

    try:
        livro = load_workbook(BytesIO(conteudo), read_only=True, data_only=True)
    except Exception as erro:
        raise ArquivoInvalido("não consegui abrir a planilha") from erro

    try:
        for aba in livro.worksheets[:ABAS_MAXIMAS]:
            linhas = list(aba.iter_rows(max_col=COLUNAS_LIDAS, values_only=True))
            encontrado = achar_cabecalho(linhas)
            if encontrado:
                break
        else:
            raise ArquivoInvalido("não achei o cabeçalho (preciso de colunas de data, produto, quantidade e preço)")
    finally:
        livro.close()

    inicio, mapa = encontrado
    if len(linhas) - inicio > LINHAS_MAXIMAS:
        raise ArquivoInvalido("mais de 50 mil linhas")

    resultado = ArquivoLido(arquivo=nome, loja=nome_da_loja(nome, posicao))
    for indice, linha in enumerate(linhas[inicio + 1 :], start=inicio + 2):
        if all(vazio(v) for v in linha):
            continue
        venda, motivo = ler_linha(linha, mapa)
        if motivo:
            resultado.descartadas.append({"arquivo": nome, "linha": indice, "motivo": motivo, "conteudo": _conteudo(linha)})
            continue
        venda["loja"] = venda["loja"] or resultado.loja
        venda["arquivo"] = nome
        venda["linha"] = indice
        resultado.vendas.append(venda)
    return resultado
